from pypdf import PdfReader
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from openpyxl import load_workbook
import pandas as pd
import os


class DocumentChunker:

    def extract_text(self, file_path):

        extension = os.path.splitext(file_path)[1].lower()

        if extension == ".pdf":
            return self.extract_pdf(file_path)

        elif extension == ".docx":
            return self.extract_docx(file_path)

        elif extension == ".txt":
            return self.extract_txt(file_path)

        elif extension == ".csv":
            return self.extract_csv(file_path)

        elif extension == ".pptx":
            return self.extract_pptx(file_path)

        elif extension == ".xlsx":
            return self.extract_xlsx(file_path)

        else:
            raise Exception(
                f"Unsupported file type: {extension}"
            )

    def extract_pdf(self, file_path):

        reader = PdfReader(file_path)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

        return text

    def extract_docx(self, file_path):

        doc = DocxDocument(file_path)

        return "\n".join(
            para.text
            for para in doc.paragraphs
        )

    def extract_txt(self, file_path):

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as file:

            return file.read()

    def extract_csv(self, file_path):

        df = pd.read_csv(file_path)

        return df.to_string()

    def extract_pptx(self, file_path):

        presentation = Presentation(file_path)

        text = ""

        for slide in presentation.slides:
            for shape in slide.shapes:

                if hasattr(shape, "text"):
                    text += shape.text + "\n"

        return text

    def extract_xlsx(self, file_path):

        workbook = load_workbook(
            file_path,
            data_only=True
        )

        text = ""

        for sheet in workbook.worksheets:

            text += f"\nSheet: {sheet.title}\n"

            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(cell)
                    for cell in row
                    if cell is not None
                ]

                text += " ".join(values) + "\n"

        return text

    def chunk_text(self, text):

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200
        )

        return splitter.split_text(text)