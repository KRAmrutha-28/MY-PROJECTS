from openpyxl import Workbook

def create_xlsx(report_text, output_path):

    wb = Workbook()

    ws = wb.active

    ws.title = "Research Report"

    ws["A1"] = "Research Report"

    row = 3

    for line in report_text.split("\n"):

        ws.cell(
            row=row,
            column=1,
            value=line
        )

        row += 1

    wb.save(output_path)

    return output_path