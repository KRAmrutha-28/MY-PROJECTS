from pptx import Presentation

def create_pptx(report_text, output_path):

    prs = Presentation()

    slide = prs.slides.add_slide(
        prs.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "Research Report"
    )

    slide.placeholders[1].text = (
        report_text[:2000]
    )

    prs.save(output_path)

    return output_path