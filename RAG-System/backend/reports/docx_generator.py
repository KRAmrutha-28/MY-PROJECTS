from docx import Document

def create_docx(report_text, output_path):

    doc = Document()

    doc.add_heading(
        "Research Report",
        level=1
    )

    doc.add_paragraph(report_text)

    doc.save(output_path)

    return output_path